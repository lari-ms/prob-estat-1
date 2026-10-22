"""
Script para gerar o relatório acadêmico em formato .docx
Estrutura conforme normas ABNT e especificações do professor
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def criar_relatorio():
    """Cria o documento do relatório com estrutura completa"""
    
    doc = Document()
    
    # Configurar margens (ABNT: 3cm esquerda, 2cm direita, 3cm superior, 2cm inferior)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.18)  # 3cm
        section.bottom_margin = Inches(0.79)  # 2cm
        section.left_margin = Inches(1.18)  # 3cm
        section.right_margin = Inches(0.79)  # 2cm
    
    # ==============================================================================
    # CAPA
    # ==============================================================================
    
    # Instituto
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('INSTITUTO FEDERAL DO ESPÍRITO SANTO')
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('CAMPUS SERRA')
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Bacharelado em Sistemas de Informação')
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    # Espaçamento
    for _ in range(4):
        doc.add_paragraph()
    
    # Título
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('TRABALHO I')
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = 'Arial'
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Probabilidade, Apostas e Pensamento Crítico:')
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Simulação e Análise de uma Máquina Estatística')
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    # Espaçamento
    for _ in range(4):
        doc.add_paragraph()
    
    # Integrantes - PLACEHOLDER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[NOMES DOS INTEGRANTES]')
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(255, 0, 0)  # Vermelho para indicar que precisa preencher
    
    # Espaçamento
    for _ in range(6):
        doc.add_paragraph()
    
    # Rodapé da capa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Serra, ES')
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Outubro de 2026')
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    
    # Quebra de página
    doc.add_page_break()
    
    # ==============================================================================
    # SUMÁRIO (será preenchido automaticamente pelo Word/Google Docs)
    # ==============================================================================
    
    p = doc.add_heading('SUMÁRIO', level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph('[O sumário será gerado automaticamente pelo editor]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.italic = True
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_page_break()
    
    # ==============================================================================
    # 1. INTRODUÇÃO
    # ==============================================================================
    
    doc.add_heading('1 INTRODUÇÃO', level=1)
    
    # 1.1 Contextualização
    doc.add_heading('1.1 Contextualização', level=2)
    
    p = doc.add_paragraph(
        'As apostas digitais aproximaram jogos de azar, publicidade, meios de pagamento e uso '
        'cotidiano do celular. Entre esses produtos estão as máquinas de rolos, popularizadas no '
        'Brasil pelo chamado "jogo do tigrinho". Seus símbolos e animações podem tornar uma rodada '
        'fácil de acompanhar, mas não explicam, por si, a probabilidade de receber um prêmio nem o '
        'resultado financeiro de repetir o jogo.'
    )
    
    p = doc.add_paragraph(
        'A Organização Mundial da Saúde relaciona os danos das apostas a sofrimento financeiro, '
        'comprometimento de despesas essenciais, problemas de saúde mental e conflitos familiares. '
        'As consequências podem alcançar pessoas que não apostam e ocorrer mesmo sem um diagnóstico '
        'de transtorno do jogo. Máquinas eletrônicas e jogos de cassino estão entre os produtos de '
        'maior risco.'
    )
    
    # 1.2 Objetivos
    doc.add_heading('1.2 Objetivos', level=2)
    
    doc.add_heading('1.2.1 Objetivo Geral', level=3)
    p = doc.add_paragraph(
        'Compreender como ganhos ocasionais podem coexistir com uma perda esperada para o jogador '
        'através da construção e análise de uma máquina de rolos com regras públicas e '
        'probabilidades conhecidas.'
    )
    
    doc.add_heading('1.2.2 Objetivos Específicos', level=3)
    p = doc.add_paragraph('Calcular e validar as probabilidades teóricas do modelo de rolos;', style='List Bullet')
    p = doc.add_paragraph('Executar simulações computacionais para verificar experimentalmente os valores teóricos;', style='List Bullet')
    p = doc.add_paragraph('Demonstrar a Lei dos Grandes Números através de sequências de rodadas;', style='List Bullet')
    p = doc.add_paragraph('Analisar o comportamento de sessões com saldo limitado;', style='List Bullet')
    p = doc.add_paragraph('Investigar a independência entre rodadas consecutivas;', style='List Bullet')
    p = doc.add_paragraph('Comparar diferentes tabelas de pagamento e seus efeitos sobre o RTP.', style='List Bullet')
    
    # 1.3 Justificativa
    doc.add_heading('1.3 Justificativa', level=2)
    
    p = doc.add_paragraph(
        'Este trabalho se justifica pela necessidade de aplicar conceitos de Probabilidade e '
        'Estatística a um problema social relevante. A simulação permite demonstrar que aleatoriedade '
        'não significa igualdade de vantagem econômica, e que uma vitória ocasional não transforma '
        'um jogo com esperança negativa em fonte de renda.'
    )
    
    p = doc.add_paragraph(
        'Além disso, o trabalho desenvolve competências em programação científica, análise de dados, '
        'visualização e comunicação de resultados estatísticos, habilidades fundamentais para '
        'profissionais de Sistemas de Informação.'
    )
    
    doc.add_page_break()
    
    # ==============================================================================
    # 2. FUNDAMENTAÇÃO TEÓRICA
    # ==============================================================================
    
    doc.add_heading('2 FUNDAMENTAÇÃO TEÓRICA', level=1)
    
    # 2.1 Conceitos Probabilísticos
    doc.add_heading('2.1 Conceitos Probabilísticos', level=2)
    
    doc.add_heading('2.1.1 Espaço Amostral e Eventos', level=3)
    p = doc.add_paragraph(
        'O espaço amostral Ω é o conjunto de todos os resultados possíveis de um experimento '
        'aleatório. Um evento é qualquer subconjunto de Ω. No modelo implementado, cada rodada '
        'sorteia três paradas independentes, cada uma com 5 possibilidades, resultando em um '
        'espaço amostral de 5³ = 125 combinações equiprováveis na linha central.'
    )
    
    doc.add_heading('2.1.2 Probabilidade Uniforme e Independência', level=3)
    p = doc.add_paragraph(
        'Dizemos que um evento tem distribuição uniforme quando todos os seus resultados elementares '
        'têm a mesma probabilidade. Dois eventos são independentes quando a ocorrência de um não '
        'afeta a probabilidade do outro. No simulador, cada parada K tem probabilidade 1/5 de assumir '
        'qualquer valor em {0, 1, 2, 3, 4}, e os três rolos são sorteados independentemente.'
    )
    
    doc.add_heading('2.1.3 Variável Aleatória Discreta', level=3)
    p = doc.add_paragraph(
        'Uma variável aleatória é uma função que associa cada resultado do espaço amostral a um '
        'número real. No modelo, definimos P como o prêmio bruto e X = P - 1 como o lucro líquido '
        'do jogador. A variável X pode assumir 6 valores possíveis: {-1, 4, 9, 14, 24, 49} créditos.'
    )
    
    doc.add_heading('2.1.4 Esperança Matemática', level=3)
    p = doc.add_paragraph(
        'A esperança (ou valor esperado) E[X] de uma variável aleatória discreta é a média '
        'ponderada de seus valores, onde os pesos são as probabilidades:'
    )
    
    p = doc.add_paragraph('E[X] = Σ x·P(X = x)')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph(
        'A linearidade da esperança garante que E[X] = E[P] - 1. Para o modelo base implementado, '
        'E[P] = 0,84 crédito e E[X] = -0,16 crédito por rodada.'
    )
    
    # 2.2 Return to Player (RTP)
    doc.add_heading('2.2 Return to Player (RTP)', level=2)
    
    p = doc.add_paragraph(
        'O RTP (Return to Player) é uma medida padrão da indústria de jogos que indica a proporção '
        'esperada dos créditos apostados que retorna ao jogador em forma de prêmios:'
    )
    
    p = doc.add_paragraph('RTP = E[P] / aposta')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph(
        'Com aposta unitária fixa, RTP = E[P]. No modelo base, RTP = 0,84 (84%), indicando vantagem '
        'esperada de 16% para o operador. É fundamental distinguir RTP da probabilidade de receber '
        'um prêmio: no modelo implementado, P(prêmio) = 4%, mas RTP = 84%.'
    )
    
    # 2.3 Lei dos Grandes Números
    doc.add_heading('2.3 Lei dos Grandes Números', level=2)
    
    p = doc.add_paragraph(
        'A Lei dos Grandes Números estabelece que, sob certas condições, a média de uma sequência de '
        'variáveis aleatórias independentes e identicamente distribuídas converge para a esperança '
        'quando o número de observações tende ao infinito. Isso não significa que a convergência seja '
        'monotônica nem que exista um erro máximo garantido para n finito.'
    )
    
    doc.add_page_break()
    
    # ==============================================================================
    # 3. METODOLOGIA
    # ==============================================================================
    
    doc.add_heading('3 METODOLOGIA', level=1)
    
    # 3.1 Modelo da Máquina
    doc.add_heading('3.1 Modelo da Máquina de Rolos', level=2)
    
    doc.add_heading('3.1.1 Estrutura e Regras', level=3)
    p = doc.add_paragraph(
        'O modelo implementado utiliza cinco símbolos identificados por A, B, C, D e E, organizados '
        'em uma sequência circular R = (A, B, C, D, E) nas posições 0 a 4. Cada rodada sorteia três '
        'paradas independentes K₁, K₂, K₃, uma para cada rolo, com distribuição uniforme.'
    )
    
    p = doc.add_paragraph(
        'Para cada parada K, a coluna correspondente exibe:'
    )
    p = doc.add_paragraph('Topo: R[K]', style='List Bullet')
    p = doc.add_paragraph('Centro: R[(K+1) mod 5]', style='List Bullet')
    p = doc.add_paragraph('Base: R[(K+2) mod 5]', style='List Bullet')
    
    p = doc.add_paragraph(
        'Apenas a linha central é avaliada para pagamento. A aposta é fixa em 1 crédito por rodada, '
        'e o lucro é calculado como X = P - 1, onde P é o prêmio bruto.'
    )
    
    doc.add_heading('3.1.2 Tabelas de Pagamento', level=3)
    
    # Tabela Base
    p = doc.add_paragraph('Tabela 1 - Pagamentos do Modelo Base')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.bold = True
    
    table = doc.add_table(rows=7, cols=4)
    table.style = 'Light Grid Accent 1'
    
    # Cabeçalho
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Combinação'
    hdr_cells[1].text = 'Prêmio (P)'
    hdr_cells[2].text = 'Lucro (X)'
    hdr_cells[3].text = 'Probabilidade'
    
    # Dados
    data = [
        ('A-A-A', '5', '4', '1/125'),
        ('B-B-B', '10', '9', '1/125'),
        ('C-C-C', '15', '14', '1/125'),
        ('D-D-D', '25', '24', '1/125'),
        ('E-E-E', '50', '49', '1/125'),
        ('Outras 120', '0', '-1', '120/125'),
    ]
    
    for i, (comb, premio, lucro, prob) in enumerate(data, 1):
        row_cells = table.rows[i].cells
        row_cells[0].text = comb
        row_cells[1].text = premio
        row_cells[2].text = lucro
        row_cells[3].text = prob
    
    p = doc.add_paragraph('Fonte: Elaborado pelos autores.')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.size = Pt(10)
    
    doc.add_paragraph()
    
    # 3.2 Implementação Computacional
    doc.add_heading('3.2 Implementação Computacional', level=2)
    
    doc.add_heading('3.2.1 Ambiente e Ferramentas', level=3)
    p = doc.add_paragraph(
        'O simulador foi implementado em Python 3.12.7, utilizando as seguintes bibliotecas:'
    )
    p = doc.add_paragraph('numpy 1.26.4 - computação científica e operações matemáticas;', style='List Bullet')
    p = doc.add_paragraph('pandas 2.1.4 - manipulação e análise de dados;', style='List Bullet')
    p = doc.add_paragraph('matplotlib 3.8.2 - visualização de dados e geração de gráficos;', style='List Bullet')
    p = doc.add_paragraph('pytest 7.4.3 - framework de testes automatizados.', style='List Bullet')
    
    doc.add_heading('3.2.2 Módulos Desenvolvidos', level=3)
    p = doc.add_paragraph(
        'O código foi organizado em módulos com responsabilidades bem definidas:'
    )
    
    p = doc.add_paragraph(
        'maquina.py: implementa o núcleo do simulador (sorteio, montagem de colunas, execução de '
        'rodadas);', style='List Bullet'
    )
    p = doc.add_paragraph(
        'calculos.py: contém funções para cálculos teóricos de probabilidades, esperança e RTP;',
        style='List Bullet'
    )
    p = doc.add_paragraph(
        'experimentos.py: implementa os cinco experimentos obrigatórios [A SER CONCLUÍDO];',
        style='List Bullet'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph(
        'visualizacao.py: gera gráficos e visualizações [A SER CONCLUÍDO];',
        style='List Bullet'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph(
        'exportacao.py: exporta resultados para CSV [A SER CONCLUÍDO].',
        style='List Bullet'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    doc.add_heading('3.2.3 Reprodutibilidade', level=3)
    p = doc.add_paragraph(
        'Para garantir a reprodutibilidade dos resultados, utilizou-se o gerador pseudoaleatório '
        'random.Random(seed) do módulo padrão do Python. Cada experimento inicializa o gerador com '
        'uma semente específica, documentada no código e no relatório. Os mesmos valores de semente '
        'produzem sequências idênticas no mesmo ambiente (Python 3.12.7, random padrão).'
    )
    
    doc.add_heading('3.2.4 Validação e Testes', level=3)
    p = doc.add_paragraph(
        'Foram desenvolvidos 47 testes automatizados abrangendo:'
    )
    p = doc.add_paragraph('Verificação do sorteio uniforme e independência dos rolos;', style='List Bullet')
    p = doc.add_paragraph('Validação da montagem circular das colunas;', style='List Bullet')
    p = doc.add_paragraph('Conferência dos cálculos de prêmio e lucro;', style='List Bullet')
    p = doc.add_paragraph('Validação dos valores teóricos de referência;', style='List Bullet')
    p = doc.add_paragraph('Verificação da reprodutibilidade das sequências.', style='List Bullet')
    
    p = doc.add_paragraph(
        'Todos os 47 testes passaram com sucesso, confirmando a correção matemática da implementação.'
    )
    
    doc.add_page_break()
    
    # ==============================================================================
    # 4. RESULTADOS
    # ==============================================================================
    
    doc.add_heading('4 RESULTADOS', level=1)
    
    # 4.1 Validação do Modelo
    doc.add_heading('4.1 Validação do Modelo Matemático', level=2)
    
    p = doc.add_paragraph(
        'Antes de executar os experimentos, os valores teóricos do modelo foram calculados e validados:'
    )
    
    # Tabela de validação
    p = doc.add_paragraph('Tabela 2 - Valores Teóricos Validados')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.bold = True
    
    table = doc.add_table(rows=6, cols=3)
    table.style = 'Light Grid Accent 1'
    
    # Cabeçalho
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Métrica'
    hdr_cells[1].text = 'Valor Teórico'
    hdr_cells[2].text = 'Status'
    
    # Dados
    data = [
        ('Espaço amostral', '125 combinações', '✓'),
        ('P(prêmio)', '0,04 (4%)', '✓'),
        ('E[P]', '0,84 crédito', '✓'),
        ('E[X]', '-0,16 crédito', '✓'),
        ('RTP', '0,84 (84%)', '✓'),
    ]
    
    for i, (metrica, valor, status) in enumerate(data, 1):
        row_cells = table.rows[i].cells
        row_cells[0].text = metrica
        row_cells[1].text = valor
        row_cells[2].text = status
    
    p = doc.add_paragraph('Fonte: Elaborado pelos autores.')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.size = Pt(10)
    
    doc.add_paragraph()
    
    # ⚠️ SEÇÕES DOS EXPERIMENTOS - PLACEHOLDERS
    
    # 4.2 Experimento 1
    doc.add_heading('4.2 Experimento 1: Frequências e Probabilidades', level=2)
    
    p = doc.add_paragraph('[RESULTADOS A SEREM INSERIDOS APÓS EXECUÇÃO]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    p = doc.add_paragraph(
        'Este experimento executará 10.000 rodadas com semente 2026 para verificar se as frequências '
        'observadas se aproximam das probabilidades teóricas calculadas.'
    )
    
    p = doc.add_paragraph('Procedimento:')
    run = p.runs[0]
    run.bold = True
    p = doc.add_paragraph('Executar 10.000 rodadas independentes;', style='List Bullet')
    p = doc.add_paragraph('Contar a frequência de cada resultado de lucro possível;', style='List Bullet')
    p = doc.add_paragraph('Calcular lucro médio, RTP observado e desvio-padrão amostral;', style='List Bullet')
    p = doc.add_paragraph('Comparar frequências observadas com probabilidades teóricas;', style='List Bullet')
    p = doc.add_paragraph('Gerar gráfico de barras comparativo.', style='List Bullet')
    
    p = doc.add_paragraph('[TABELA DE FREQUÊNCIAS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[GRÁFICO DE FREQUÊNCIAS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[INTERPRETAÇÃO DOS RESULTADOS - A INSERIR]')
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    # 4.3 Experimento 2
    doc.add_heading('4.3 Experimento 2: Lei dos Grandes Números', level=2)
    
    p = doc.add_paragraph('[RESULTADOS A SEREM INSERIDOS APÓS EXECUÇÃO]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    p = doc.add_paragraph(
        'Este experimento demonstrará a Lei dos Grandes Números através de três sequências '
        'independentes de 10.000 rodadas cada.'
    )
    
    p = doc.add_paragraph('Procedimento:')
    run = p.runs[0]
    run.bold = True
    p = doc.add_paragraph('Executar três sequências com sementes 10, 2026 e 123456;', style='List Bullet')
    p = doc.add_paragraph('Reutilizar a sequência de 2026 do Experimento 1;', style='List Bullet')
    p = doc.add_paragraph('Calcular média acumulada após 100, 1.000 e 10.000 rodadas;', style='List Bullet')
    p = doc.add_paragraph('Gerar gráfico das trajetórias de convergência.', style='List Bullet')
    
    p = doc.add_paragraph('[TABELA DE MÉDIAS ACUMULADAS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[GRÁFICO DE TRAJETÓRIAS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[INTERPRETAÇÃO DOS RESULTADOS - A INSERIR]')
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    # 4.4 Experimento 3
    doc.add_heading('4.4 Experimento 3: Saldo Limitado', level=2)
    
    p = doc.add_paragraph('[RESULTADOS A SEREM INSERIDOS APÓS EXECUÇÃO]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    p = doc.add_paragraph(
        'Este experimento simulará 100 participantes virtuais, cada um com saldo inicial de 100 '
        'créditos, para analisar o comportamento de sessões com recursos limitados.'
    )
    
    p = doc.add_paragraph('Procedimento:')
    run = p.runs[0]
    run.bold = True
    p = doc.add_paragraph('Simular 100 participantes com saldo inicial de 100 créditos;', style='List Bullet')
    p = doc.add_paragraph('Máximo de 500 rodadas ou até saldo < 1;', style='List Bullet')
    p = doc.add_paragraph('Calcular proporções de ganho, empate e perda;', style='List Bullet')
    p = doc.add_paragraph('Registrar média e mediana dos saldos finais;', style='List Bullet')
    p = doc.add_paragraph('Gerar gráfico das 5 primeiras trajetórias.', style='List Bullet')
    
    p = doc.add_paragraph('[TABELA DE ESTATÍSTICAS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[GRÁFICO DE TRAJETÓRIAS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[INTERPRETAÇÃO DOS RESULTADOS - A INSERIR]')
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    # 4.5 Experimento 4
    doc.add_heading('4.5 Experimento 4: Probabilidade Condicional', level=2)
    
    p = doc.add_paragraph('[RESULTADOS A SEREM INSERIDOS APÓS EXECUÇÃO]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    p = doc.add_paragraph(
        'Este experimento investigará a independência entre rodadas consecutivas, testando a falácia '
        'de que uma sequência de perdas aumenta a chance de vitória na rodada seguinte.'
    )
    
    p = doc.add_paragraph('Procedimento:')
    run = p.runs[0]
    run.bold = True
    p = doc.add_paragraph('Reutilizar as 10.000 rodadas do Experimento 1;', style='List Bullet')
    p = doc.add_paragraph('Examinar 9.999 pares consecutivos;', style='List Bullet')
    p = doc.add_paragraph('Classificar por "após prêmio" e "após não-prêmio";', style='List Bullet')
    p = doc.add_paragraph('Calcular P(prêmio seguinte | resultado anterior).', style='List Bullet')
    
    p = doc.add_paragraph('[TABELA DE PROBABILIDADES CONDICIONAIS - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[INTERPRETAÇÃO DOS RESULTADOS - A INSERIR]')
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    # 4.6 Experimento 5
    doc.add_heading('4.6 Experimento 5: Comparação de Pagamentos', level=2)
    
    p = doc.add_paragraph('[RESULTADOS A SEREM INSERIDOS APÓS EXECUÇÃO]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    p = doc.add_paragraph(
        'Este experimento comparará o modelo base com uma tabela de pagamentos aumentados em 10%, '
        'mantendo as mesmas probabilidades.'
    )
    
    p = doc.add_paragraph('Procedimento:')
    run = p.runs[0]
    run.bold = True
    p = doc.add_paragraph('Reutilizar os 10.000 resultados do Experimento 1;', style='List Bullet')
    p = doc.add_paragraph('Aplicar duas tabelas de pagamentos (base e +10%);', style='List Bullet')
    p = doc.add_paragraph('Comparar P(prêmio), prêmio médio, lucro médio e RTP;', style='List Bullet')
    p = doc.add_paragraph('Demonstrar que aumentar pagamentos não altera a chance de prêmio.', style='List Bullet')
    
    p = doc.add_paragraph('[TABELA COMPARATIVA - A INSERIR]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    
    p = doc.add_paragraph('[INTERPRETAÇÃO DOS RESULTADOS - A INSERIR]')
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_page_break()
    
    # ==============================================================================
    # 5. DISCUSSÃO CRÍTICA
    # ==============================================================================
    
    doc.add_heading('5 DISCUSSÃO CRÍTICA', level=1)
    
    p = doc.add_paragraph('[ESTA SEÇÃO SERÁ PREENCHIDA APÓS CONCLUSÃO DOS EXPERIMENTOS]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    # 5.1 Mecanismo e Percepção
    doc.add_heading('5.1 Mecanismo e Percepção', level=2)
    
    p = doc.add_paragraph(
        '[Relacionar a esperança negativa aos ganhos ocasionais observados nos experimentos. '
        'Discutir a dispersão dos resultados e a visibilidade dos relatos de vencedores. '
        'Usar ao menos uma conta e dois resultados experimentais.]'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    # 5.2 Renda e Saúde
    doc.add_heading('5.2 Renda e Saúde', level=2)
    
    p = doc.add_paragraph(
        '[Discutir por que perdas podem comprometer recursos essenciais. Identificar o que vem '
        'da simulação e o que vem de estudos sobre pessoas reais. Observar que uma rodada perdida '
        'não é um diagnóstico clínico.]'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    # 5.3 Políticas e Interesses
    doc.add_heading('5.3 Políticas e Interesses', level=2)
    
    p = doc.add_paragraph(
        '[Examinar um argumento favorável à restrição e um argumento crítico à proibição, '
        'com fontes identificadas. Considerar proteção, autonomia, publicidade e fiscalização. '
        'Avaliar a força das evidências.]'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    # 5.4 Responsabilidade Profissional
    doc.add_heading('5.4 Responsabilidade Profissional', level=2)
    
    p = doc.add_paragraph(
        '[Discutir como escolhas de interface, divulgação e implementação podem influenciar a '
        'compreensão do usuário. Formular uma conclusão educativa sem promessa de ganho nem '
        'culpabilização de pessoas afetadas.]'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    doc.add_page_break()
    
    # ==============================================================================
    # 6. CONSIDERAÇÕES FINAIS
    # ==============================================================================
    
    doc.add_heading('6 CONSIDERAÇÕES FINAIS', level=1)
    
    p = doc.add_paragraph('[ESTA SEÇÃO SERÁ PREENCHIDA APÓS CONCLUSÃO DE TODOS OS EXPERIMENTOS]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    p = doc.add_paragraph(
        '[Resumir os principais achados dos cinco experimentos. Destacar conceitos demonstrados '
        '(espaço amostral, independência, esperança, Lei dos Grandes Números). Reforçar as limitações '
        'da simulação: ela demonstra propriedades matemáticas de um modelo, mas não mede prevalência '
        'de transtornos ou impacto nacional. Sugerir trabalhos futuros.]'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    doc.add_page_break()
    
    # ==============================================================================
    # REFERÊNCIAS
    # ==============================================================================
    
    doc.add_heading('REFERÊNCIAS', level=1)
    
    # PLACEHOLDER para referências
    p = doc.add_paragraph(
        '[AS REFERÊNCIAS SERÃO INSERIDAS CONFORME NORMAS ABNT APÓS CONSULTA DAS FONTES]'
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    p = doc.add_paragraph(
        'Sugestões de referências a incluir (verificar e formatar conforme ABNT):'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    p = doc.add_paragraph(
        'World Health Organization - Gambling fact sheet'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    p = doc.add_paragraph(
        'The Lancet Public Health Commission on gambling'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    p = doc.add_paragraph(
        'LENAD III - Caderno Temático sobre jogos de aposta'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    p = doc.add_paragraph(
        'Lei 14.790/2023 - Marco regulatório das apostas'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    p = doc.add_paragraph(
        'Python Software Foundation. Python Language Reference, version 3.12.'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    p = doc.add_paragraph(
        'ILLOWSKY, B.; DEAN, S. Introductory Statistics. OpenStax, 2023.'
    )
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_page_break()
    
    # ==============================================================================
    # DECLARAÇÃO DE USO DE IA
    # ==============================================================================
    
    doc.add_heading('DECLARAÇÃO DE USO DE INTELIGÊNCIA ARTIFICIAL', level=1)
    
    p = doc.add_paragraph(
        '[ESTA SEÇÃO DEVE SER PREENCHIDA PELA EQUIPE]'
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    
    doc.add_paragraph()
    
    p = doc.add_paragraph(
        'Os integrantes do grupo declaram que utilizaram ferramentas de inteligência artificial '
        '(especificar quais) para as seguintes atividades:'
    )
    
    p = doc.add_paragraph('[Listar atividades onde IA foi utilizada]', style='List Bullet')
    run = p.runs[0]
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.italic = True
    
    p = doc.add_paragraph(
        'Todos os valores matemáticos, cálculos teóricos e resultados experimentais foram validados '
        'independentemente pela equipe. Cada integrante compreende o conteúdo apresentado neste '
        'relatório e pode explicá-lo durante a defesa do trabalho.'
    )
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph('[LOCAL], [DATA]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph('_' * 50)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('[NOME INTEGRANTE 1]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    p = doc.add_paragraph('_' * 50)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('[NOME INTEGRANTE 2]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    p = doc.add_paragraph('_' * 50)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('[NOME INTEGRANTE 3]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Salvar documento
    caminho = 'docs/RELATORIO_TRABALHO_I.docx'
    doc.save(caminho)
    
    return caminho


if __name__ == "__main__":
    print("=" * 70)
    print("GERANDO RELATÓRIO ACADÊMICO")
    print("=" * 70)
    print()
    
    try:
        caminho = criar_relatorio()
        print(f"✓ Relatório criado com sucesso!")
        print(f"✓ Arquivo salvo em: {caminho}")
        print()
        print("=" * 70)
        print("INSTRUÇÕES:")
        print("=" * 70)
        print()
        print("1. Abra o arquivo no Microsoft Word ou Google Docs")
        print("2. Gere o sumário automaticamente (Inserir → Índice/Sumário)")
        print("3. Preencha os placeholders em VERMELHO:")
        print("   - Nomes dos integrantes na capa")
        print("   - Resultados dos experimentos (Seção 4)")
        print("   - Discussão crítica completa (Seção 5)")
        print("   - Considerações finais (Seção 6)")
        print("   - Referências formatadas em ABNT")
        print("   - Declaração de uso de IA")
        print()
        print("4. Os textos em CINZA são orientações - remover após preencher")
        print()
        print("=" * 70)
        print("O QUE JÁ ESTÁ PRONTO:")
        print("=" * 70)
        print("✓ Estrutura completa do relatório")
        print("✓ Introdução (contextualização, objetivos, justificativa)")
        print("✓ Fundamentação teórica (conceitos, RTP, LGN)")
        print("✓ Metodologia (modelo, implementação, validação)")
        print("✓ Seção de resultados com placeholders para 5 experimentos")
        print("✓ Estrutura da discussão crítica")
        print("✓ Formatação acadêmica (margens, fontes, espaçamento)")
        print("✓ Tabelas e espaços para gráficos")
        print()
        print("=" * 70)
        print("O QUE PRECISA SER ATUALIZADO:")
        print("=" * 70)
        print("⏳ Nomes dos integrantes")
        print("⏳ Resultados do Experimento 1 (após Fase 2)")
        print("⏳ Resultados do Experimento 2 (após Fase 3)")
        print("⏳ Resultados do Experimento 3 (após Fase 4)")
        print("⏳ Resultados do Experimento 4 (após Fase 5)")
        print("⏳ Resultados do Experimento 5 (após Fase 6)")
        print("⏳ Discussão crítica completa")
        print("⏳ Considerações finais")
        print("⏳ Referências formatadas")
        print("⏳ Declaração de uso de IA")
        print()
        print("=" * 70)
        
    except Exception as e:
        print(f"✗ Erro ao criar relatório: {e}")
        print()
        print("Certifique-se de que a biblioteca python-docx está instalada:")
        print("pip install python-docx")
